"""Administrator test cases, immutable execution snapshots and human reviews."""
from __future__ import annotations

import base64
import binascii
import csv
from datetime import datetime
from io import BytesIO, StringIO
import re
from typing import Literal
import uuid
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import ParseError

from pydantic import BaseModel, ConfigDict, Field, ValidationError, field_validator
from sqlalchemy import ForeignKey, Integer, JSON, String, Text
from sqlalchemy.orm import Mapped, mapped_column

from .db import Base, UTCDateTime
from .models import _utcnow

MODULE = Literal["module_1", "module_2", "module_3", "module_4"]
HEADERS = ["Case ID", "模块", "用户类型", "测试场景", "用户起始输入", "AI应达到的目标"]
FIELDS = ["case_code", "module", "user_type", "scenario", "user_input", "expected_behavior"]


class EvaluationCase(Base):
    __tablename__ = "evaluation_cases"
    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=lambda: str(uuid.uuid4()))
    case_code: Mapped[str] = mapped_column(String(64), unique=True)
    module: Mapped[str] = mapped_column(String(16))
    user_type: Mapped[str] = mapped_column(Text, default="")
    scenario: Mapped[str] = mapped_column(Text, default="")
    user_input: Mapped[str] = mapped_column(Text)
    expected_behavior: Mapped[str] = mapped_column(Text, default="")
    extra_columns: Mapped[dict] = mapped_column(JSON, default=dict)
    revision: Mapped[int] = mapped_column(Integer, default=1)
    created_by: Mapped[int] = mapped_column(ForeignKey("user_accounts.id"))
    updated_at: Mapped[datetime] = mapped_column(UTCDateTime(), default=_utcnow)


class EvaluationRun(Base):
    __tablename__ = "evaluation_runs"
    id: Mapped[str] = mapped_column(String(36), primary_key=True)
    case_id: Mapped[str] = mapped_column(ForeignKey("evaluation_cases.id"), index=True)
    parent_run_id: Mapped[str | None] = mapped_column(ForeignKey("evaluation_runs.id"), nullable=True)
    created_by: Mapped[int] = mapped_column(ForeignKey("user_accounts.id"))
    status: Mapped[str] = mapped_column(String(16), default="running")
    case_snapshot: Mapped[dict] = mapped_column(JSON)
    prompt_snapshot: Mapped[dict] = mapped_column(JSON)
    provider: Mapped[str] = mapped_column(String(32))
    model: Mapped[str | None] = mapped_column(String(128), nullable=True)
    input_text: Mapped[str] = mapped_column(Text)
    reply: Mapped[str] = mapped_column(Text, default="")
    transcript: Mapped[list] = mapped_column(JSON, default=list)
    memory: Mapped[dict] = mapped_column(JSON, default=dict)
    metrics: Mapped[dict] = mapped_column(JSON, default=dict)
    error_code: Mapped[str | None] = mapped_column(String(80), nullable=True)
    created_at: Mapped[datetime] = mapped_column(UTCDateTime(), default=_utcnow)
    finished_at: Mapped[datetime | None] = mapped_column(UTCDateTime(), nullable=True)


class EvaluationReview(Base):
    __tablename__ = "evaluation_reviews"
    id: Mapped[str] = mapped_column(String(36), primary_key=True, default=lambda: str(uuid.uuid4()))
    run_id: Mapped[str] = mapped_column(ForeignKey("evaluation_runs.id"), index=True)
    reviewer_id: Mapped[int] = mapped_column(ForeignKey("user_accounts.id"))
    verdict: Mapped[str] = mapped_column(String(16))
    notes: Mapped[str] = mapped_column(Text, default="")
    created_at: Mapped[datetime] = mapped_column(UTCDateTime(), default=_utcnow)


class CaseInput(BaseModel):
    model_config = ConfigDict(extra="forbid", str_strip_whitespace=True)
    case_code: str = Field(min_length=1, max_length=64)
    module: MODULE
    user_type: str = Field(default="", max_length=4000)
    scenario: str = Field(default="", max_length=6000)
    user_input: str = Field(min_length=1, max_length=12000)
    expected_behavior: str = Field(default="", max_length=12000)
    extra_columns: dict[str, str] = Field(default_factory=dict, max_length=40)

    @field_validator("extra_columns")
    @classmethod
    def check_extra(cls, value):
        if any(len(k) > 120 or len(v) > 12000 for k, v in value.items()):
            raise ValueError("额外列过长")
        return value


class CaseUpdate(CaseInput):
    revision: int = Field(ge=1)


class ImportRequest(BaseModel):
    filename: str = Field(max_length=255)
    content_base64: str = Field(max_length=7_000_000)
    sheet: str | None = Field(default=None, max_length=100)


class ImportCommit(BaseModel):
    cases: list[CaseInput] = Field(min_length=1, max_length=500)


class RunRequest(BaseModel):
    request_id: uuid.UUID
    provider: Literal["claude", "deepseek", "doubao"]
    parent_run_id: uuid.UUID | None = None
    message: str | None = Field(default=None, min_length=1, max_length=12000)


class ReviewInput(BaseModel):
    verdict: Literal["pass", "fail", "unreviewed"]
    notes: str = Field(default="", max_length=12000)


def case_data(row: EvaluationCase) -> dict:
    return {"id": row.id, **{key: getattr(row, key) for key in FIELDS},
            "extra_columns": row.extra_columns, "revision": row.revision,
            "updated_at": row.updated_at.isoformat()}


def run_data(row: EvaluationRun) -> dict:
    # A process restart can leave an unfinished record. Expose it, never silently retry/bill twice.
    status = row.status
    if status in ("queued", "running") and (_utcnow() - row.created_at).total_seconds() > 240:
        status = "interrupted"
    return {"id": row.id, "case_id": row.case_id, "parent_run_id": row.parent_run_id,
            "status": status, "case_snapshot": row.case_snapshot, "provider": row.provider,
            "model": row.model, "input_text": row.input_text, "reply": row.reply,
            "transcript": row.transcript, "metrics": row.metrics, "error_code": row.error_code,
            "created_at": row.created_at.isoformat(),
            "finished_at": row.finished_at.isoformat() if row.finished_at else None}


def _header(value: str) -> str:
    return re.sub(r"[\s_]+", "", value).lower()


def parse_upload(payload: ImportRequest) -> dict:
    """Bounded values-only import. Never execute formulas or fetch external links."""
    try:
        raw = base64.b64decode(payload.content_base64, validate=True)
    except (ValueError, binascii.Error) as exc:
        raise ValueError("文件编码无效") from exc
    if len(raw) > 5_000_000:
        raise ValueError("文件不能超过 5 MB")
    sheets = []
    if payload.filename.lower().endswith(".csv"):
        try:
            rows = list(csv.reader(StringIO(raw.decode("utf-8-sig"))))
        except (UnicodeError, csv.Error) as exc:
            raise ValueError("CSV 请另存为 UTF-8 编码") from exc
        sheets, selected = ["CSV"], "CSV"
    elif payload.filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook
        try:
            with ZipFile(BytesIO(raw)) as archive:
                if sum(info.file_size for info in archive.infolist()) > 25_000_000:
                    raise ValueError("工作簿解压后超过 25 MB")
            book = load_workbook(BytesIO(raw), read_only=True, data_only=False, keep_links=False)
            try:
                sheets = book.sheetnames
                selected = payload.sheet or ("评测集" if "评测集" in sheets else sheets[0])
                if selected not in sheets:
                    raise ValueError("找不到所选工作表")
                sheet = book[selected]
                if sheet.max_row and sheet.max_row > 550 or sheet.max_column and sheet.max_column > 50:
                    raise ValueError("工作表最多 550 行、50 列，请只保留测试表区域")
                rows = []
                for cells in sheet.iter_rows():
                    if len(rows) >= 550 or len(cells) > 50:
                        raise ValueError("工作表最多 550 行、50 列")
                    if any(cell.data_type == "f" for cell in cells):
                        raise ValueError("导入表包含公式，请先复制为值后导入")
                    rows.append(["" if c.value is None else str(c.value) for c in cells])
            finally:
                book.close()
        except (BadZipFile, KeyError, OSError, ParseError, IndexError) as exc:
            raise ValueError("无法读取 XLSX，请确认文件格式") from exc
        except ValueError:
            raise
        except Exception as exc:
            # XML engines differ by installed optional dependencies (e.g. lxml).
            raise ValueError("无法读取 XLSX，请确认文件格式") from exc
    else:
        raise ValueError("支持 .xlsx 和 UTF-8 .csv；旧 .xls 请另存为 .xlsx")
    if len(rows) > 550 or any(len(row) > 50 for row in rows):
        raise ValueError("最多 550 行、50 列")
    header_index = next((i for i, row in enumerate(rows[:20]) if "caseid" in [_header(str(v)) for v in row]), None)
    if header_index is None:
        raise ValueError("前 20 行未找到 Case ID 表头")
    headers = [str(v).strip() for v in rows[header_index]]
    positions = {}
    for index, label in enumerate(headers):
        normalized = _header(label)
        for field, expected in zip(FIELDS, HEADERS):
            if normalized == _header(expected) or (field == "user_type" and normalized.startswith("用户类型")):
                if field in positions:
                    raise ValueError(f"表头重复：{expected}")
                positions[field] = index
    missing = [label for field, label in zip(FIELDS, HEADERS) if field not in positions]
    if missing:
        raise ValueError("缺少列：" + "、".join(missing))
    extra_headers = [h for i, h in enumerate(headers) if h and i not in positions.values()]
    if len(extra_headers) != len(set(extra_headers)):
        raise ValueError("额外列表头重复，请使用不同列名")
    cases, errors, seen = [], [], set()
    for number, row in enumerate(rows[header_index + 1:], header_index + 2):
        if not any(str(v).strip() for v in row):
            continue
        value = lambda index: str(row[index]).strip() if index < len(row) else ""
        data = {field: value(index) for field, index in positions.items()}
        module = data["module"].upper()
        data["module"] = {f"M{i}": f"module_{i}" for i in range(1, 5)}.get(module, data["module"].lower())
        data["extra_columns"] = {h: value(i) for i, h in enumerate(headers) if h and i not in positions.values()}
        try:
            item = CaseInput(**data)
            if item.case_code.casefold() in seen:
                raise ValueError("文件中 Case ID 重复")
            seen.add(item.case_code.casefold())
            cases.append(item.model_dump())
        except ValidationError as exc:
            messages = []
            labels = dict(zip(FIELDS, HEADERS))
            for error in exc.errors(include_url=False, include_input=False):
                field = str(error["loc"][0])
                label = labels.get(field, field)
                if field == "module":
                    messages.append(f"模块“{data['module'] or '空'}”暂不支持自动运行（当前支持 M1～M4）")
                elif error["type"] == "string_too_short":
                    messages.append(f"{label}不能为空")
                else:
                    messages.append(f"{label}格式不符合要求")
            errors.append({"row": number, "message": "；".join(messages)})
        except ValueError as exc:
            errors.append({"row": number, "message": str(exc)[:400]})
    if len(cases) > 500:
        raise ValueError("一次最多导入 500 条用例")
    return {"sheets": sheets, "selected_sheet": selected, "cases": cases, "errors": errors}


def export_csv(rows: list[list]) -> bytes:
    output = StringIO(newline="")
    writer = csv.writer(output)
    for row in rows:
        # Spreadsheet apps must not interpret imported prompts/answers as formulas.
        writer.writerow([("'" + str(v)) if str(v).lstrip().startswith(("=", "+", "-", "@")) else str(v) for v in row])
    return output.getvalue().encode("utf-8-sig")
