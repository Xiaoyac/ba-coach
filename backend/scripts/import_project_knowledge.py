"""Import the curated ``KnowledgeBase`` directory into the shared KB tables.

The filename manifest is deliberately explicit: a new document must be
assigned a clinical category before it can become visible to the coach. This
prevents an unrelated file copied into the directory from silently entering
every user's prompt.
"""

from __future__ import annotations

import asyncio
import sys
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.db import dispose_db, get_sessionmaker, init_db  # noqa: E402
from app.knowledge_store import (  # noqa: E402
    KNOWLEDGE_CATEGORY_MODULES,
    import_knowledge_source,
)


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_KNOWLEDGE_DIR = PROJECT_ROOT / "KnowledgeBase"

SOURCE_CATEGORIES: dict[str, str] = {
    "1~6章(2).md": "BA",
    "7~9章(2).md": "BA",
    "2024-adult-compendium身体活动分类(1)(1).xlsx": "PA",
    "动机访谈法.md": "MI",
    "行为改变技术分类法第一版(1).md": "BCT",
    "运动和心理健康相关文献_3.md": "PA",
    "运动与心理健康相关文献_1(1).md": "PA",
    "运动与心理健康相关文献_2(1).md": "PA",
    "运动与心理健康相关文献_4.md": "PA",
    "知识条目1(1).md": "BCT",
}

def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def compendium_to_markdown(path: Path) -> str:
    """Convert the four-column Adult Compendium sheet into searchable text."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = tuple(_cell_text(value) for value in next(rows))
        if len(headers) < 4 or headers[0] != "Major Heading":
            raise ValueError(f"unexpected Compendium columns in {path.name}: {headers}")

        output = [
            "# 2024 Adult Compendium 身体活动分类",
            "",
            "本资料用于核对候选活动是否属于身体活动，并提供活动代码、代谢当量（MET）与英文活动描述。MET仅作分类参考，不替代个体安全评估。",
        ]
        current_heading = ""
        record_count = 0
        for raw in rows:
            values = tuple(_cell_text(value) for value in raw[:4])
            if not any(values):
                continue
            major, code, met, description = values
            if major != current_heading:
                output.extend(("", f"# 身体活动分类：{major or '未分类'}"))
                current_heading = major
            if code.isdigit() and len(code) < 5:
                code = code.zfill(5)
            output.extend(
                (
                    "",
                    f"- 活动代码：{code}",
                    f"- 代谢当量（MET）：{met}",
                    f"- 活动描述：{description}",
                )
            )
            record_count += 1
        if record_count < 1000:
            raise ValueError(
                f"Compendium import looks incomplete: only {record_count} records"
            )
        return "\n".join(output).strip() + "\n"
    finally:
        workbook.close()


def _read_source(path: Path) -> str:
    if path.suffix.lower() == ".xlsx":
        return compendium_to_markdown(path)
    return path.read_text(encoding="utf-8-sig")


def _source_paths(directory: Path) -> Iterable[tuple[Path, str]]:
    missing = [name for name in SOURCE_CATEGORIES if not (directory / name).is_file()]
    if missing:
        raise FileNotFoundError("missing knowledge source(s): " + ", ".join(missing))
    supported = {
        path.name
        for path in directory.iterdir()
        if path.is_file() and path.suffix.lower() in {".md", ".xlsx"}
    }
    unknown = sorted(supported - set(SOURCE_CATEGORIES) - {"README.md"})
    if unknown:
        raise ValueError(
            "knowledge source(s) have no category assignment: " + ", ".join(unknown)
        )
    for name, category in SOURCE_CATEGORIES.items():
        yield directory / name, category


async def run(directory: Path = DEFAULT_KNOWLEDGE_DIR) -> None:
    if not directory.is_dir():
        raise FileNotFoundError(f"knowledge directory not found: {directory}")
    await init_db()
    try:
        async with get_sessionmaker()() as db:
            for path, category in _source_paths(directory):
                result = await import_knowledge_source(
                    db,
                    name=path.name,
                    category=category,
                    markdown=_read_source(path),
                    updated_by="project-knowledge-import",
                )
                modules = ",".join(KNOWLEDGE_CATEGORY_MODULES[category])
                state = "unchanged" if result.unchanged else "imported"
                print(
                    f"{path.name}: {category} -> {modules}; "
                    f"{result.chunk_count} chunks ({state})"
                )
            await db.commit()
    finally:
        await dispose_db()


if __name__ == "__main__":
    asyncio.run(run())
